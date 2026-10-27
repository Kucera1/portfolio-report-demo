"""Generátor SYNTETICKÝCH dat fiktivního e-shopu s kávou „Ukázková pražírna“.

Vše je vymyšlené a vygenerované s pevným seedem (stejný výsledek při každém
spuštění). Data jsou záměrně „neuklizená“, jak to v malých firmách bývá:
různé formáty dat, překlepy v kategoriích a kódech, prázdné řádky,
mezisoučty uprostřed tabulky, duplicitní řádky z opakovaného exportu.

Spuštění:  python -m report_demo.generator  [cílová složka]
"""

import csv
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from .nacitani import SOUBOR_NAKLADU, SOUBOR_OBJEDNAVEK, SOUBOR_SKLADU

SEED = 2026
ZACATEK = date(2026, 8, 3)   # pondělí
KONEC = date(2026, 10, 11)   # neděle
VYPADEK_ESHOPU = date(2026, 10, 6)      # den s minimem objednávek
FIREMNI_OBJEDNAVKA = datetime(2026, 10, 8, 10, 12)
KAMPAN = date(2026, 10, 7)

# kód, název, prodejní cena, nákupní cena, oblíbenost, skladem, minimální zásoba
PRODUKTY = [
    ("KAV-ETI-250", "Etiopie Yirgacheffe 250 g", 289, 150, 0.16, 60, 80),
    ("KAV-ETI-1000", "Etiopie Yirgacheffe 1 kg", 990, 520, 0.05, 120, 30),
    ("KAV-KOL-250", "Kolumbie Huila 250 g", 259, 135.5, 0.14, 400, 60),
    ("KAV-KOL-1000", "Kolumbie Huila 1 kg", 899, 470, 0.05, 140, 30),
    ("KAV-BRA-250", "Brazílie Santos 250 g", 219, 110, 0.13, 380, 60),
    ("KAV-BRA-1000", "Brazílie Santos 1 kg", 749, 380, 0.07, 200, 30),
    ("KAV-KEN-250", "Keňa AA 250 g", 319, 170, 0.08, 900, 40),     # předzásobeno – točí se pomalu
    ("KAV-ESP-1000", "Espresso směs Demo 1 kg", 799, 400, 0.12, 330, 50),
    ("KAV-DEC-250", "Bezkofeinová Mexiko 250 g", 279, 145, 0.06, 160, 30),
    ("PRI-V60", "Dripper V60 keramický", 549, 260, 0.04, 110, 20),
    ("PRI-MLY-RUC", "Ruční mlýnek na kávu", 1290, 690, 0.02, 4, 6),
    ("PRI-FIL-100", "Papírové filtry 100 ks", 129, 45, 0.08, 45, 60),
]

# Doplňkové zboží, které se prodává jen občas (pro přehled pomalu se točícího zboží).
# kód, název, prodejní cena, nákupní cena, skladem, minimální zásoba
DOPLNKOVE_PRODUKTY = [
    ("PRI-SAL-200", "Keramický šálek 200 ml", 249, 110, 140, 10),
    ("PRI-SIR-KAR", "Kávový sirup karamel 250 ml", 189, 150, 60, 10),
    ("PRI-KRA-DAR", "Dárková krabička", 99, 40, 120, 20),
]
# Šálek se kupuje hlavně spolu s dripperem – z objednávek pak jde poznat vhodný balíček.
SALEK_K_DRIPPERU = 0.35
# Doplňkové zboží se prodávalo jen do 6. 9.; týdny 37–41 (a tedy čísla v ukázkovém
# reportu za týden 41 a jeho srovnání) zůstávají beze změny.
DOPLNKY_OD = date(2026, 8, 17)
DOPLNKY_DO = date(2026, 9, 7)        # bez tohoto dne
SIRUP_DNY = [date(2026, 8, 20), date(2026, 8, 31)]
KRABICKA_DNY = [date(2026, 8, 5)]

# Náklady na balení a vyřízení jednoho kusu (Kč). Platební poplatek v % je v nastavení sortimentu.
POPLATKY_KS = {
    "KAV-ETI-250": 12, "KAV-ETI-1000": 18, "KAV-KOL-250": 12, "KAV-KOL-1000": 18,
    "KAV-BRA-250": 12, "KAV-BRA-1000": 18, "KAV-KEN-250": 12, "KAV-ESP-1000": 18,
    "KAV-DEC-250": 12, "PRI-V60": 20, "PRI-MLY-RUC": 25, "PRI-FIL-100": 8,
    "PRI-SAL-200": 25, "PRI-SIR-KAR": 15, "PRI-KRA-DAR": 6,
}

VARIANTY_KATEGORII = {
    "Doprava": ["Doprava", "doprava", "Dopravaa", " Doprava ", "Poštovné"],
    "Marketing": ["Marketing", "marketing", "Marketng", "Reklama"],
    "Obaly": ["Obaly", "obaly", "Obal"],
    "Nájem": ["Nájem", "najem"],
    "Energie": ["Energie", "Enrgie"],
    "Mzdy": ["Mzdy", "mzdy"],
    "Software": ["Software", "SW"],
}
MESICE = {8: "srpen", 9: "září", 10: "říjen"}


def _dny():
    d = ZACATEK
    while d <= KONEC:
        yield d
        d += timedelta(days=1)


def _cena_text(cena: float) -> str:
    return f"{cena:.2f}".replace(".", ",")


def vygeneruj_objednavky(rng: np.random.Generator) -> list[list[str]]:
    vahy = np.array([p[4] for p in PRODUKTY])
    radky = []
    cislo = 1
    for den in _dny():
        intenzita = 22 * (1.25 if den.weekday() >= 5 else 1.0) * (1 + 0.004 * (den - ZACATEK).days)
        if den == VYPADEK_ESHOPU:
            intenzita *= 0.08
        for _ in range(rng.poisson(intenzita)):
            cas = datetime(den.year, den.month, den.day, int(rng.integers(7, 23)), int(rng.integers(0, 60)))
            pocet_polozek = int(rng.choice([1, 2, 3], p=[0.6, 0.3, 0.1]))
            for i in rng.choice(len(PRODUKTY), size=pocet_polozek, replace=False, p=vahy):
                kod, nazev, cena = PRODUKTY[i][:3]
                stav = "Stornováno" if rng.random() < 0.04 else "Vyřízeno"
                radky.append([f"26{cislo:05d}", cas, kod, nazev, int(rng.choice([1, 1, 1, 2, 2, 3])), cena, stav])
            cislo += 1
        if den == FIREMNI_OBJEDNAVKA.date():
            radky.append([f"26{cislo:05d}", FIREMNI_OBJEDNAVKA, "KAV-ESP-1000",
                          "Espresso směs Demo 1 kg", 60, 719, "Vyřízeno"])
            cislo += 1

    # „Nepořádek“ typický pro exporty a ruční zásahy.
    neupravene = []
    for cislo_obj, cas, kod, nazev, ks, cena, stav in radky:
        nahoda = rng.random()
        if kod == "KAV-ETI-250" and nahoda < 0.08:
            kod = "KAV-ETI-25O"          # písmeno O místo nuly
        elif kod == "KAV-BRA-250" and nahoda < 0.05:
            kod = "KAV-BRA250"
        elif nahoda < 0.05:
            kod = kod.lower()
        elif nahoda < 0.08:
            kod = f" {kod} "
        stav = str(rng.choice({"Vyřízeno": ["Vyřízeno"] * 12 + ["vyřízeno", "VYRIZENO", "Vyrizeno "],
                                "Stornováno": ["Stornováno", "storno"]}[stav]))
        datum_text = cas.strftime("%Y-%m-%d %H:%M:%S") if rng.random() < 0.02 else cas.strftime("%d.%m.%Y %H:%M")
        mnozstvi = "" if rng.random() < 0.004 else str(ks)
        neupravene.append([cislo_obj, datum_text, kod, nazev, mnozstvi, _cena_text(cena), stav])

    # Opakovaný export: některé řádky jsou v souboru dvakrát za sebou.
    # Doplňkové položky se přidají za řádek objednávky, ke které patří (vlastní generátor
    # náhodných čísel, takže se ostatní data nezmění).
    doplnky = _doplnkove_polozky(radky, neupravene)
    vysledek = []
    for i, radek in enumerate(neupravene):
        vysledek.append(radek)
        if rng.random() < 0.015:
            vysledek.append(list(radek))
        vysledek.extend(doplnky.get(i, []))
    return vysledek


def _doplnkove_polozky(radky: list, neupravene: list) -> dict[int, list[list[str]]]:
    rng = np.random.default_rng(SEED + 1)
    doplnky: dict[int, list[list[str]]] = {}
    prvni_v_den: dict[date, int] = {}
    for i, (_, cas, kod, _, _, _, stav) in enumerate(radky):
        if stav == "Vyřízeno":
            prvni_v_den.setdefault(cas.date(), i)

    def pridej(i: int, kod: str) -> None:
        _, nazev, cena = next(p for p in DOPLNKOVE_PRODUKTY if p[0] == kod)[:3]
        cislo_obj, datum_text = neupravene[i][:2]
        doplnky.setdefault(i, []).append([cislo_obj, datum_text, kod, nazev, "1", _cena_text(cena), "Vyřízeno"])

    for i, (_, cas, kod, _, _, _, stav) in enumerate(radky):
        if (kod == "PRI-V60" and stav == "Vyřízeno" and DOPLNKY_OD <= cas.date() < DOPLNKY_DO
                and rng.random() < SALEK_K_DRIPPERU):
            pridej(i, "PRI-SAL-200")
    for den in SIRUP_DNY:
        pridej(prvni_v_den[den], "PRI-SIR-KAR")
    for den in KRABICKA_DNY:
        pridej(prvni_v_den[den], "PRI-KRA-DAR")
    return doplnky


def vygeneruj_naklady(rng: np.random.Generator) -> list[tuple[date, str, str, float]]:
    polozky = []
    pondeli = ZACATEK
    while pondeli <= KONEC:
        for posun in (1, 4):
            polozky.append((pondeli + timedelta(days=posun), "Doprava", "Svoz zásilek – dopravce",
                            round(float(rng.normal(4200, 500)))))
        polozky.append((pondeli + timedelta(days=2), "Obaly", "Krabice a výplň",
                        round(float(rng.normal(2100, 300)), 2)))
        polozky.append((pondeli + timedelta(days=3), "Marketing", "Reklama ve vyhledávači",
                        round(float(rng.normal(5000, 700)))))
        pondeli += timedelta(days=7)
    polozky.append((KAMPAN, "Marketing", "Kampaň na sociálních sítích – podzimní akce", 22000))
    for mesic in (8, 9, 10):
        polozky += [
            (date(2026, mesic, 1), "Nájem", "Nájem sklad a kancelář", 32000),
            (date(2026, mesic, 5), "Software", "Předplatné e-shop platformy", 1490),
            (date(2026, mesic, 10), "Mzdy", "Mzdy – brigádníci balení", 58000),
            (date(2026, mesic, 15), "Energie", "Elektřina a plyn", round(float(rng.normal(3900, 200)))),
        ]
    return sorted(p for p in polozky if ZACATEK <= p[0] <= KONEC)


def _zapis_datum(bunka, den: date, rng: np.random.Generator) -> None:
    volba = rng.random()
    if volba < 0.4:
        bunka.value = datetime(den.year, den.month, den.day)
        bunka.number_format = "DD.MM.YYYY"
    elif volba < 0.6:
        bunka.value = f"{den.day}.{den.month}.{den.year}"
    elif volba < 0.75:
        bunka.value = den.isoformat()
    elif volba < 0.9:
        bunka.value = den.strftime("%d/%m/%Y")
    else:
        bunka.value = f"{den.day}. {den.month}. {den.year}"


def _castka_hodnota(castka: float, rng: np.random.Generator):
    volba = rng.random()
    tisice = f"{castka:,.0f}".replace(",", " ")
    if volba < 0.5:
        return castka
    if volba < 0.7:
        return f"{tisice} Kč"
    if volba < 0.85:
        return f"{castka:.2f}".replace(".", ",")
    return f"{castka:,.2f} Kč".replace(",", " ").replace(".", ",")


def uloz_naklady(cesta: Path, rng: np.random.Generator) -> None:
    polozky = vygeneruj_naklady(rng)
    # Duplicitně zapsané položky (někdo je přepsal z faktury podruhé).
    for i in rng.choice(len(polozky), size=3, replace=False):
        polozky.insert(int(i) + 1, polozky[int(i)])

    wb = Workbook()
    ws = wb.active
    ws.title = "Náklady"
    ws.append(["Ukázková pražírna – náklady 2026 (pracovní tabulka)"])
    ws.append(["zapisuje provoz, průběžně"])
    ws.append([])
    ws.append(["Datum", "Kategorie ", "Popis", "Částka"])
    mesic = polozky[0][0].month
    soucet = 0.0
    for den, kategorie, popis, castka in polozky:
        if den.month != mesic:
            ws.append(["", "", f"Celkem {MESICE[mesic]}", round(soucet, 2)])
            mesic, soucet = den.month, 0.0
        if rng.random() < 0.12:
            ws.append([])
        ws.append([None, str(rng.choice(VARIANTY_KATEGORII[kategorie])), popis, _castka_hodnota(castka, rng)])
        _zapis_datum(ws.cell(row=ws.max_row, column=1), den, rng)
        soucet += castka
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 46
    ws.column_dimensions["D"].width = 16
    wb.properties.creator = "generator syntetickych dat"
    wb.properties.created = datetime(2026, 10, 12)
    wb.save(cesta)


def uloz_sklad(cesta: Path) -> None:
    radky = []
    polozky = [(kod, nazev, cena, nakup, skladem, minimum)
               for kod, nazev, cena, nakup, _, skladem, minimum in PRODUKTY] + DOPLNKOVE_PRODUKTY
    for kod, nazev, cena, nakup, skladem, minimum in polozky:
        radky.append([kod, nazev, str(skladem), str(minimum), _cena_text(nakup).removesuffix(",00"),
                      str(cena), str(POPLATKY_KS[kod])])
    radky[4][0] = " kav-bra-250"          # malá písmena a mezera
    radky[9][2] = "-2"                    # chyba inventury
    radky[6][3] = ""                      # chybí minimální zásoba
    radky.insert(2, ["KAV-KOL-250", "Kolumbie Huila 250 g", "380", "60", "135,50", "259", "12"])  # starší řádek
    with open(cesta, "w", newline="", encoding="utf-8") as f:
        zapis = csv.writer(f)
        zapis.writerow(["sku", "nazev", "skladem_ks", "min_zasoba", "nakupni_cena", "prodejni_cena", "poplatek_ks"])
        zapis.writerows(radky)


def vygeneruj_vse(slozka: Path) -> Path:
    slozka = Path(slozka)
    slozka.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(SEED)
    with open(slozka / SOUBOR_OBJEDNAVEK, "w", newline="", encoding="utf-8-sig") as f:
        zapis = csv.writer(f, delimiter=";")
        zapis.writerow(["Číslo objednávky", "Datum", "Kód zboží", "Název zboží", "Množství", "Cena/ks", "Stav"])
        zapis.writerows(vygeneruj_objednavky(rng))
    uloz_naklady(slozka / SOUBOR_NAKLADU, rng)
    uloz_sklad(slozka / SOUBOR_SKLADU)
    return slozka


if __name__ == "__main__":
    cil = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "data" / "vstupy"
    print(f"Syntetická data uložena do: {vygeneruj_vse(cil)}")
