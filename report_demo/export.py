"""Export reportu do Excelu (pro další práci s čísly) a PDF (pro čtení a tisk)."""

import io
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from fpdf import FPDF  # noqa: E402
from fpdf.fonts import FontFace  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from . import formaty  # noqa: E402
from .report import TydenniReport, formatuj_hodnotu, formatuj_zmenu, souhrnna_tabulka  # noqa: E402

NAZEV_FIRMY = "Ukázková pražírna"
MODRA = "#2a78d6"
SEDA = "#8a8a86"
CERVENA = "#d03b3b"

F_KC = '#,##0 "Kč"'
F_KS = "#,##0"
F_DESETINY = "#,##0.0"
F_PCT = "0.0 %"
F_ZMENA = "+0.0 %;-0.0 %;0.0 %"
F_KC_ZNAMENKO = '+#,##0 "Kč";-#,##0 "Kč";0 "Kč"'
F_DATUM = "d. m. yyyy"
FORMATY_SLOUPCU = {
    "Tržby": F_KC, "Skutečnost": F_KC, "Obvyklá úroveň": F_KC, "Tento týden": F_KC,
    "Průměr za týden (předchozí 4 týdny)": F_KC, "Prodáno ks": F_KS, "Skladem ks": F_KS,
    "Minimální zásoba": F_KS, "Doporučeno objednat ks": F_KS, "Počet": F_KS,
    "Prodej za den (ø 4 týdny)": F_DESETINY, "Vystačí na dní": F_DESETINY,
    "Podíl na tržbách": F_PCT, "Změna proti minulému týdnu": F_ZMENA, "Odchylka": F_ZMENA,
    "Den": F_DATUM,
}
FORMATY_UKAZATELU = {"kc": F_KC, "ks": F_KS, "pct": F_PCT}


def nazev_souboru(report: TydenniReport, pripona: str) -> str:
    rok, tyden, _ = report.od.isocalendar()
    return f"tydenni_report_{rok}-T{tyden:02d}.{pripona}"


# ---------------------------------------------------------------- Excel

def _uprav_list(ws, zacatek: int = 1) -> None:
    hlavicka = ws[zacatek]
    for bunka in hlavicka:
        bunka.font = Font(bold=True)
        bunka.fill = PatternFill("solid", fgColor="E8EEF7")
        bunka.alignment = Alignment(wrap_text=True, vertical="top")
        format_cisla = FORMATY_SLOUPCU.get(bunka.value)
        for radek in range(zacatek + 1, ws.max_row + 1):
            hodnota = ws.cell(row=radek, column=bunka.column)
            if format_cisla:
                hodnota.number_format = format_cisla
            if isinstance(hodnota.value, float) and (math.isnan(hodnota.value) or math.isinf(hodnota.value)):
                hodnota.value = None
    for sloupec in ws.columns:
        delka = max((len(str(b.value)) for b in sloupec[zacatek - 1:] if b.value is not None), default=8)
        ws.column_dimensions[sloupec[0].column_letter].width = min(max(delka + 2, 10), 70)
    ws.freeze_panes = ws.cell(row=zacatek + 1, column=1)


def export_excel(report: TydenniReport, cesta: Path) -> Path:
    souhrn = souhrnna_tabulka(report)
    listy = {
        "Top produkty": report.top_produkty,
        "Nízká zásoba": report.nizka_zasoba,
        "Výkyvy": report.vykyvy,
        "Denní tržby": report.denni_trzby,
        "Náklady podle kategorií": report.naklady_kategorie,
        "Úpravy dat (souhrn)": report.log_souhrn,
        "Úpravy dat (detail)": report.log_detail,
    }
    with pd.ExcelWriter(cesta, engine="openpyxl") as zapis:
        souhrn.drop(columns="typ").to_excel(zapis, sheet_name="Souhrn", index=False, startrow=3)
        for nazev, tabulka in listy.items():
            tabulka.to_excel(zapis, sheet_name=nazev, index=False)

        ws = zapis.sheets["Souhrn"]
        ws["A1"] = f"Týdenní report – {NAZEV_FIRMY}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = formaty.popis_tydne(report.od)
        _uprav_list(ws, zacatek=4)
        for i, typ in enumerate(souhrn["typ"], start=5):
            ws.cell(row=i, column=2).number_format = FORMATY_UKAZATELU[typ]
            ws.cell(row=i, column=3).number_format = FORMATY_UKAZATELU[typ]
            ws.cell(row=i, column=4).number_format = '+0.0" p. b.";-0.0" p. b."' if typ == "pct" else F_ZMENA
            if typ == "pct" and ws.cell(row=i, column=4).value is not None:
                ws.cell(row=i, column=4).value *= 100
        radek = ws.max_row + 2
        ws.cell(row=radek, column=1, value="Co ovlivnilo provozní zisk").font = Font(bold=True)
        ws.cell(row=radek + 1, column=1, value=report.veta_rozkladu)
        radek += 2
        for sloupec, nazev in enumerate(report.rozklad_zisku.columns, start=1):
            bunka = ws.cell(row=radek, column=sloupec, value=nazev)
            bunka.font = Font(bold=True)
            bunka.fill = PatternFill("solid", fgColor="E8EEF7")
        for vliv, dopad, vysvetleni in report.rozklad_zisku.itertuples(index=False):
            radek += 1
            ws.cell(row=radek, column=1, value=vliv)
            ws.cell(row=radek, column=2, value=round(dopad, 2)).number_format = F_KC_ZNAMENKO
            ws.cell(row=radek, column=3, value=vysvetleni)
        radek += 2
        ws.cell(row=radek, column=1, value="Hlavní zjištění").font = Font(bold=True)
        for veta in report.zjisteni + [""] + report.poznamky:
            radek += 1
            ws.cell(row=radek, column=1, value=veta)
        ws.column_dimensions["A"].width = 44
        for nazev in listy:
            _uprav_list(zapis.sheets[nazev])
    return cesta


# ---------------------------------------------------------------- PDF

def _cesta_k_pismu(soubor: str) -> str:
    return str(Path(matplotlib.get_data_path()) / "fonts" / "ttf" / soubor)


def graf_dennich_trzeb(report: TydenniReport) -> bytes:
    denni = report.denni_trzby
    vykyvove_dny = set(report.vykyvy["Kdy"])
    popisky = [formaty.datum_kratce(d) for d in denni["Den"]]
    barvy = [CERVENA if p in vykyvove_dny else MODRA for p in popisky]

    fig, ax = plt.subplots(figsize=(7.2, 2.6), dpi=160)
    ax.bar(popisky, denni["Tržby"], color=barvy, width=0.6, label="Tržby")
    ax.plot(popisky, denni["Obvyklá úroveň"], linestyle="none", marker="_", markersize=22,
            markeredgewidth=2, color=SEDA, label="Obvyklá úroveň")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: formaty.cislo(v)))
    ax.set_ylabel("Kč")
    for strana in ("top", "right"):
        ax.spines[strana].set_visible(False)
    ax.grid(axis="y", color="#e5e5e2", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


class _Pdf(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("DejaVu", size=7)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, f"Ukázka se syntetickými daty – fiktivní firma · strana {self.page_no()}/{{nb}}",
                  align="C")


def _nadpis(pdf: FPDF, text: str, min_vyska: float = 45) -> None:
    if pdf.will_page_break(min_vyska):  # nadpis nenechávat osamocený na konci stránky
        pdf.add_page()
    pdf.ln(3)
    pdf.set_font("DejaVu", style="B", size=12)
    pdf.cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", size=9)


def _tabulka(pdf: FPDF, hlavicka: list[str], radky: list[list[str]], sirky: tuple, zarovnani: tuple) -> None:
    styl = FontFace(emphasis="BOLD", fill_color=(232, 238, 247))
    with pdf.table(col_widths=sirky, text_align=zarovnani, headings_style=styl, line_height=5.5,
                   borders_layout="HORIZONTAL_LINES") as tabulka:
        for radek in [hlavicka] + radky:
            r = tabulka.row()
            for hodnota in radek:
                r.cell(hodnota)


def export_pdf(report: TydenniReport, cesta: Path) -> Path:
    pdf = _Pdf(format="A4")
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_font("DejaVu", fname=_cesta_k_pismu("DejaVuSans.ttf"))
    pdf.add_font("DejaVu", style="B", fname=_cesta_k_pismu("DejaVuSans-Bold.ttf"))
    pdf.add_page()

    pdf.set_font("DejaVu", style="B", size=18)
    pdf.cell(0, 10, f"Týdenní report – {NAZEV_FIRMY}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", size=10)
    pdf.cell(0, 6, formaty.popis_tydne(report.od), new_x="LMARGIN", new_y="NEXT")

    _nadpis(pdf, "Hlavní zjištění")
    for veta in report.zjisteni:
        pdf.multi_cell(0, 5, f"• {veta}", new_x="LMARGIN", new_y="NEXT")

    _nadpis(pdf, "Souhrn")
    souhrn = souhrnna_tabulka(report)
    _tabulka(pdf, ["Ukazatel", "Tento týden", "Minulý týden", "Změna"],
             [[r.Ukazatel, formatuj_hodnotu(r[1], r.typ), formatuj_hodnotu(r[2], r.typ), formatuj_zmenu(r.Změna, r.typ)]
              for r in souhrn.itertuples(index=False)],
             (70, 40, 40, 30), ("LEFT", "RIGHT", "RIGHT", "RIGHT"))

    _nadpis(pdf, "Co ovlivnilo provozní zisk")
    pdf.multi_cell(0, 5, report.veta_rozkladu, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    _tabulka(pdf, ["Vliv", "Dopad na zisk", "Vysvětlení"],
             [[v, formaty.kc_znamenko(d), t] for v, d, t in report.rozklad_zisku.itertuples(index=False)],
             (55, 30, 95), ("LEFT", "RIGHT", "LEFT"))

    _nadpis(pdf, "Tržby po dnech", min_vyska=85)  # graf je vysoký asi 65 mm
    pdf.image(io.BytesIO(graf_dennich_trzeb(report)), w=180)
    pdf.set_font("DejaVu", size=7)
    pdf.multi_cell(0, 4, "Šedá čárka = obvyklá úroveň (medián stejného dne v týdnu za předchozí 4 týdny). "
                         "Červeně dny s nezvyklým výkyvem.", new_x="LMARGIN", new_y="NEXT")

    _nadpis(pdf, "Nejprodávanější produkty")
    _tabulka(pdf, ["Produkt", "Prodáno ks", "Tržby", "Podíl", "Změna"],
             [[r.Produkt, formaty.cislo(r[2]), formaty.kc(r.Tržby), formaty.procento(r[4]),
               formaty.procento(r[5], znamenko=True)] for r in report.top_produkty.itertuples(index=False)],
             (70, 25, 35, 25, 25), ("LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"))

    _nadpis(pdf, f"Zboží k doobjednání (sklad k {formaty.datum(report.sklad_ke_dni)})")
    if report.nizka_zasoba.empty:
        pdf.cell(0, 6, "Všechny položky mají dostatečnou zásobu.", new_x="LMARGIN", new_y="NEXT")
    else:
        _tabulka(pdf, ["Produkt", "Skladem", "Minimum", "Vystačí na dní", "Objednat ks"],
                 [[r.Produkt, formaty.cislo(r[2]), formaty.cislo(r[3]),
                   "–" if math.isinf(r[5]) else formaty.cislo(r[5], 1), formaty.cislo(r[6])]
                  for r in report.nizka_zasoba.itertuples(index=False)],
                 (70, 25, 25, 30, 30), ("LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"))

    _nadpis(pdf, "Nezvyklé výkyvy")
    if report.vykyvy.empty:
        pdf.cell(0, 6, "Tento týden bez nezvyklých výkyvů.", new_x="LMARGIN", new_y="NEXT")
    else:
        _tabulka(pdf, ["Co", "Kdy", "Skutečnost", "Obvykle", "Poznámka"],
                 [[r.Co, r.Kdy, formaty.kc(r.Skutečnost), formaty.kc(r[3]), r.Poznámka]
                  for r in report.vykyvy.itertuples(index=False)],
                 (40, 20, 25, 25, 70), ("LEFT", "LEFT", "RIGHT", "RIGHT", "LEFT"))

    _nadpis(pdf, "Poznámky k datům")
    pdf.set_font("DejaVu", size=8)
    for veta in report.poznamky:
        pdf.multi_cell(0, 4.5, f"• {veta}", new_x="LMARGIN", new_y="NEXT")

    pdf.output(str(cesta))
    return cesta
