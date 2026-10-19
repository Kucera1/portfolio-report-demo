"""Schvalovací krok, exporty a celý průchod nad vygenerovanými syntetickými daty."""

from datetime import datetime

import pytest
from openpyxl import load_workbook

from report_demo.generator import vygeneruj_vse
from report_demo.nacitani import nacti_a_vycisti
from report_demo.report import dostupne_tydny, spocitej_report
from report_demo.schvaleni import (STAV_NAVRH, STAV_ODESLANO, STAV_ZAHOZENO, nacti_stav, schval_a_odesli,
                                   uloz_navrh, zahod_navrh)


def _report(tmp_path):
    data = nacti_a_vycisti(vygeneruj_vse(tmp_path / "vstupy"))
    return data, spocitej_report(data, dostupne_tydny(data.objednavky)[0])


def _soubory(slozka):
    return sorted(p.name for p in slozka.rglob("*") if p.is_file()) if slozka.exists() else []


def test_generator_je_deterministicky(tmp_path):
    prvni = vygeneruj_vse(tmp_path / "a")
    druhy = vygeneruj_vse(tmp_path / "b")
    for soubor in ("objednavky.csv", "sklad.csv"):
        assert (prvni / soubor).read_bytes() == (druhy / soubor).read_bytes()
    hodnoty = [list(load_workbook(s / "naklady.xlsx").active.values) for s in (prvni, druhy)]
    assert hodnoty[0] == hodnoty[1]


def test_report_na_syntetickych_datech_najde_vlozene_udalosti(tmp_path):
    data, report = _report(tmp_path)
    assert report.od.isoformat() == "2026-10-05"
    assert set(report.vykyvy["Co"]) == {"Tržby nezvykle nízké", "Tržby nezvykle vysoké",
                                        "Náklady nezvykle vysoké: Marketing"}
    assert {"PRI-V60", "PRI-MLY-RUC", "KAV-ETI-250", "PRI-FIL-100"} == set(report.nizka_zasoba["Kód"])
    # v nákladech nesmí zůstat mezisoučty ani neznámé kategorie
    assert not data.naklady["popis"].str.contains("Celkem").any()
    assert "Ostatní" not in set(data.naklady["kategorie"])
    assert data.objednavky.duplicated().sum() == 0


def test_rozklad_zisku_na_syntetickych_datech(tmp_path):
    _, report = _report(tmp_path)
    u, m = report.ukazatele, report.ukazatele_minule
    assert u["mesicni_platby"] == 58000 + 1490                # mzdy 10. 10. a software 5. 10.
    assert m["mesicni_platby"] == 32000                       # nájem 1. 10.
    dopady = report.rozklad_zisku.set_index("Vliv")["Dopad na zisk"]
    assert dopady.iloc[:3].sum() == pytest.approx(dopady.iloc[3])
    assert dopady.iloc[3] == pytest.approx(u["provozni_zisk"] - m["provozni_zisk"])
    assert "mzdy 58\u00a0000\u00a0Kč" in report.veta_rozkladu
    assert u["zisk_bez_mesicnich"] == pytest.approx(u["provozni_zisk"] + u["mesicni_platby"])


def test_navrh_se_bez_schvaleni_neodesle(tmp_path):
    _, report = _report(tmp_path)
    odeslano = tmp_path / "odeslano"
    navrh = uloz_navrh(report, tmp_path / "navrhy")
    assert nacti_stav(navrh)["stav"] == STAV_NAVRH
    assert _soubory(navrh) == ["navrh.json", "tydenni_report_2026-T41.pdf", "tydenni_report_2026-T41.xlsx"]
    assert _soubory(odeslano) == []


def test_schvaleni_odesle_prave_jednou(tmp_path):
    _, report = _report(tmp_path)
    odeslano = tmp_path / "odeslano"
    navrh = uloz_navrh(report, tmp_path / "navrhy", ted=datetime(2026, 10, 12, 8, 0))
    cil = schval_a_odesli(navrh, odeslano, ted=datetime(2026, 10, 12, 8, 30))
    assert nacti_stav(navrh)["stav"] == STAV_ODESLANO
    assert _soubory(cil) == ["tydenni_report_2026-T41.pdf", "tydenni_report_2026-T41.xlsx", "zprava.txt"]
    zprava = (cil / "zprava.txt").read_text(encoding="utf-8")
    assert "SIMULACE" in zprava and "týden 41/2026" in zprava
    assert len((odeslano / "odeslane_reporty.csv").read_text(encoding="utf-8-sig").splitlines()) == 2
    with pytest.raises(ValueError):
        schval_a_odesli(navrh, odeslano)


def test_zahozeny_navrh_nejde_odeslat(tmp_path):
    _, report = _report(tmp_path)
    navrh = uloz_navrh(report, tmp_path / "navrhy")
    zahod_navrh(navrh)
    assert nacti_stav(navrh)["stav"] == STAV_ZAHOZENO
    with pytest.raises(ValueError):
        schval_a_odesli(navrh, tmp_path / "odeslano")
    assert _soubory(tmp_path / "odeslano") == []


def test_dva_navrhy_ve_stejne_sekunde_se_neprepisuji(tmp_path):
    _, report = _report(tmp_path)
    ted = datetime(2026, 10, 12, 8, 0)
    prvni = uloz_navrh(report, tmp_path / "navrhy", ted=ted)
    druhy = uloz_navrh(report, tmp_path / "navrhy", ted=ted)
    assert prvni != druhy and prvni.exists() and druhy.exists()


def test_excel_obsahuje_vsechny_listy_a_cisla(tmp_path):
    _, report = _report(tmp_path)
    navrh = uloz_navrh(report, tmp_path / "navrhy")
    wb = load_workbook(navrh / "tydenni_report_2026-T41.xlsx")
    assert wb.sheetnames == ["Souhrn", "Top produkty", "Nízká zásoba", "Výkyvy", "Denní tržby",
                             "Náklady podle kategorií", "Úpravy dat (souhrn)", "Úpravy dat (detail)"]
    souhrn = wb["Souhrn"]
    assert souhrn["A5"].value == "Tržby (bez DPH)"
    assert souhrn["B5"].value == pytest.approx(report.ukazatele["trzby"])
    assert souhrn["B5"].number_format == '#,##0 "Kč"'
    assert wb["Úpravy dat (detail)"].max_row == len(report.log_detail) + 1
    hodnoty = [c.value for c in souhrn["A"]]
    assert "Provozní zisk bez měsíčních plateb" in hodnoty
    assert "Co ovlivnilo provozní zisk" in hodnoty and report.veta_rozkladu in hodnoty


def test_pdf_je_platny_soubor(tmp_path):
    _, report = _report(tmp_path)
    navrh = uloz_navrh(report, tmp_path / "navrhy")
    obsah = (navrh / "tydenni_report_2026-T41.pdf").read_bytes()
    assert obsah.startswith(b"%PDF") and len(obsah) > 10_000
